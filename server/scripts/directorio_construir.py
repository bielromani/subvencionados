"""Genera el directorio de ayuntamientos y responsables de cultura a partir de las fuentes oficiales.

Uso: python3 directorio_construir.py datos

Entrada: datos/referencia/municipios-ine.json (nomenclátor) y datos/directorio/fuentes/ (las descarga
directorio_descargar.py). Salida en datos/directorio/:
- ayuntamientos.csv: un municipio por fila, con contacto del ayuntamiento, alcalde y responsable de cultura;
- cultura.csv: concejales y consejeros cuya área incluye cultura;
- directorio-cultura.xlsx: lo anterior, con un resumen por comunidad y provincia y las fuentes.
"""
import ast
import csv
import io
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

csv.field_size_limit(10**9)

# Provincias del INE (CPRO) con su comunidad autónoma (como server/src/referencias.ts).
PROVINCIAS = {
    "01": ("Araba/Álava", "País Vasco"), "02": ("Albacete", "Castilla-La Mancha"),
    "03": ("Alicante/Alacant", "Comunitat Valenciana"), "04": ("Almería", "Andalucía"),
    "05": ("Ávila", "Castilla y León"), "06": ("Badajoz", "Extremadura"),
    "07": ("Illes Balears", "Illes Balears"), "08": ("Barcelona", "Cataluña"),
    "09": ("Burgos", "Castilla y León"), "10": ("Cáceres", "Extremadura"),
    "11": ("Cádiz", "Andalucía"), "12": ("Castellón/Castelló", "Comunitat Valenciana"),
    "13": ("Ciudad Real", "Castilla-La Mancha"), "14": ("Córdoba", "Andalucía"),
    "15": ("A Coruña", "Galicia"), "16": ("Cuenca", "Castilla-La Mancha"),
    "17": ("Girona", "Cataluña"), "18": ("Granada", "Andalucía"),
    "19": ("Guadalajara", "Castilla-La Mancha"), "20": ("Gipuzkoa", "País Vasco"),
    "21": ("Huelva", "Andalucía"), "22": ("Huesca", "Aragón"),
    "23": ("Jaén", "Andalucía"), "24": ("León", "Castilla y León"),
    "25": ("Lleida", "Cataluña"), "26": ("La Rioja", "La Rioja"),
    "27": ("Lugo", "Galicia"), "28": ("Madrid", "Comunidad de Madrid"),
    "29": ("Málaga", "Andalucía"), "30": ("Murcia", "Región de Murcia"),
    "31": ("Navarra", "Comunidad Foral de Navarra"), "32": ("Ourense", "Galicia"),
    "33": ("Asturias", "Principado de Asturias"), "34": ("Palencia", "Castilla y León"),
    "35": ("Las Palmas", "Canarias"), "36": ("Pontevedra", "Galicia"),
    "37": ("Salamanca", "Castilla y León"), "38": ("Santa Cruz de Tenerife", "Canarias"),
    "39": ("Cantabria", "Cantabria"), "40": ("Segovia", "Castilla y León"),
    "41": ("Sevilla", "Andalucía"), "42": ("Soria", "Castilla y León"),
    "43": ("Tarragona", "Cataluña"), "44": ("Teruel", "Aragón"),
    "45": ("Toledo", "Castilla-La Mancha"), "46": ("Valencia/València", "Comunitat Valenciana"),
    "47": ("Valladolid", "Castilla y León"), "48": ("Bizkaia", "País Vasco"),
    "49": ("Zamora", "Castilla y León"), "50": ("Zaragoza", "Aragón"),
    "51": ("Ceuta", "Ceuta"), "52": ("Melilla", "Melilla"),
}
ALIAS_PROVINCIA = {
    "la coruna": "15", "coruna": "15", "orense": "32", "lerida": "25", "gerona": "17", "guipuzcoa": "20",
    "vizcaya": "48", "alava": "01", "baleares": "07", "castellon": "12", "castello": "12", "valencia": "46",
    "alicante": "03", "alacant": "03",
}
ARTICULOS = "el|la|los|las|l|lo|els|les|a|o|os|as|es|sa|ses|s"
# «Cultura», «Patrimoni Cultural», «Sociocultural», «Kultura»; no «Agricultura» ni «Puericultura».
CULTURA = re.compile(r"\b(socio-?)?(cultur|kultur)", re.I)
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")


def normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFD", texto.lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[’'`´·.\-]", " ", t)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9,() ]", " ", t)).strip()


def claves(nombre: str) -> set[str]:
    """Variantes de un nombre de municipio: cada parte de "Castelló/Castellón" y el artículo pospuesto o entre paréntesis."""
    resultado = set()
    for parte in re.split(r"/", nombre):
        p = normalizar(parte)
        if not p:
            continue
        m = re.match(rf"^(.*?)\s*\(({ARTICULOS})\)$", p) or re.match(rf"^(.*), ({ARTICULOS})$", p)
        if m:
            resultado.add(f"{m[2]} {m[1]}".strip())
        sin = p.replace(",", " ").replace("(", " ").replace(")", " ")
        resultado.add(re.sub(r"\s+", " ", sin).strip())
    return resultado


def sin_articulo(clave: str) -> str:
    """"o porto do son" -> "porto do son": último recurso cuando la fuente añade o quita el artículo."""
    return re.sub(rf"^({ARTICULOS}) ", "", clave)


# Nombres de las fuentes que no coinciden con el nomenclátor del INE (provincia, nombre normalizado) -> nombre INE.
ALIAS_MUNICIPIO = {
    ("27", "alfoz do castrodouro"): "Alfoz", ("36", "cangas de morrazo"): "Cangas",
    ("32", "o castro de caldelas"): "Castro Caldelas",
}


def codigo_provincia(nombre: str) -> str | None:
    n = normalizar(nombre).replace("provincia ", "").replace("provincia de ", "")
    for cpro, (prov, _) in PROVINCIAS.items():
        if n in claves(prov):
            return cpro
    for alias, cpro in ALIAS_PROVINCIA.items():
        if n == alias or n.startswith(alias + " ") or n.endswith(" " + alias):
            return cpro
    return None


def telefono(valor) -> str:
    if valor is None:
        return ""
    t = str(valor).strip()
    if re.fullmatch(r"\d+\.0", t):
        t = t[:-2]
    numeros = re.findall(r"\d[\d\s\-.]{7,}\d", t)
    salida = []
    for n in numeros:
        d = re.sub(r"\D", "", n)
        if len(d) == 9:
            salida.append(f"{d[:3]} {d[3:6]} {d[6:]}")
    return " / ".join(dict.fromkeys(salida))


def emails(valor) -> str:
    return "; ".join(dict.fromkeys(e.lower().rstrip(".") for e in EMAIL.findall(str(valor or ""))))


def web(valor) -> str:
    w = str(valor or "").strip().split()[0] if str(valor or "").strip() else ""
    if not w or w.lower() in ("no", "no consta", "-"):
        return ""
    if not re.match(r"https?://", w, re.I):
        w = "https://" + w
    m = re.match(r"(https?://[^/]+)(.*)", w, re.I)
    return m[1].lower() + m[2] if m else w


def nombre_propio(valor: str) -> str:
    """Pone en formato de nombre los nombres escritos en mayúsculas ("JUAN DE LA FUENTE" -> "Juan de la Fuente")."""
    v = re.sub(r"\s*\((ind|independent|independiente)\.?\)", "", (valor or "").replace('"', ""), flags=re.I)
    v = re.sub(r"\s+", " ", v).strip()
    if v and re.sub(r" (i|y|de|del|la) ", " ", v) == re.sub(r" (i|y|de|del|la) ", " ", v).upper():
        v = " ".join(p if p in ("de", "del", "la", "las", "los", "y", "i") else p.capitalize()
                     for p in v.lower().split(" "))
    return v


def tokens_persona(nombre: str) -> frozenset[str]:
    """Palabras de un nombre sin apodos ni partido entre paréntesis, para reconocer a la misma persona en dos fuentes."""
    n = normalizar(re.sub(r"\(.*?\)", " ", nombre))
    return frozenset(p for p in n.replace(",", " ").split() if p not in ("i", "y", "de", "del", "la"))


def misma_persona(a: frozenset[str], b: frozenset[str]) -> bool:
    return len(a & b) >= 2 and (a <= b or b <= a or len(a & b) >= 3)


def tipo_entidad(nombre: str) -> str:
    for prefijo, tipo in (("Ajuntament", "Ayuntamiento"), ("Consell Comarcal", "Consejo comarcal"), ("Diputació", "Diputación"),
                          ("Consell General", "Conselh Generau d'Aran")):
        if nombre.startswith(prefijo):
            return tipo
    return "Otro ente local"


def leer_csv(ruta: Path, separador: str | None = None) -> list[dict]:
    crudo = ruta.read_bytes()
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("latin-1")
    if separador is None:
        primera = texto.split("\n", 1)[0]
        separador = max([",", ";", "|", "\t"], key=primera.count)
    return list(csv.DictReader(io.StringIO(texto), delimiter=separador))


class Municipios:
    def __init__(self, ruta: Path):
        self.lista = json.loads(ruta.read_text(encoding="utf-8"))
        self.por_codigo = {m["codigo"]: m for m in self.lista}
        self.indice: dict[tuple[str, str], set[str]] = defaultdict(set)
        self.indice_sin_articulo: dict[tuple[str, str], set[str]] = defaultdict(set)
        for m in self.lista:
            for k in claves(m["nombre"]):
                self.indice[(m["codigo"][:2], k)].add(m["codigo"])
                if sin_articulo(k) != k:
                    self.indice_sin_articulo[(m["codigo"][:2], sin_articulo(k))].add(m["codigo"])

    def buscar(self, nombre: str, cpro: str | None) -> str | None:
        if not cpro:
            return None
        nombre = ALIAS_MUNICIPIO.get((cpro, normalizar(nombre)), nombre)
        # Primero el nombre tal cual (con su artículo) y, si no basta, sin artículo.
        for nivel, variantes in enumerate((claves(nombre), {sin_articulo(k) for k in claves(nombre)})):
            encontrados = set()
            for k in variantes:
                encontrados |= self.indice.get((cpro, k), set())
                if nivel:
                    encontrados |= self.indice_sin_articulo.get((cpro, k), set())
            if len(encontrados) == 1:
                return next(iter(encontrados))
            if encontrados:
                return None
        return None


def main() -> None:
    datos = Path(sys.argv[1])
    fuentes = datos / "directorio" / "fuentes"
    salida = datos / "directorio"
    municipios = Municipios(datos / "referencia" / "municipios-ine.json")
    estado = json.loads((fuentes / "estado.json").read_text(encoding="utf-8"))
    descripcion = {f["fichero"]: f["descripcion"].split(" · ")[0] for f in estado["fuentes"]}

    # Ficha de cada municipio: campo -> (valor, fichero de origen). La primera fuente que da un dato gana.
    ficha: dict[str, dict[str, tuple[str, str]]] = defaultdict(dict)
    sin_ubicar: Counter = Counter()

    def anotar(codigo: str | None, fichero: str, nombre_fuente: str = "", **campos) -> None:
        if not codigo or codigo not in municipios.por_codigo:
            sin_ubicar[fichero] += 1
            return
        for campo, valor in campos.items():
            valor = (valor or "").strip() if isinstance(valor, str) else valor
            if valor and campo not in ficha[codigo]:
                ficha[codigo][campo] = (valor, fichero)

    def existe(fichero: str) -> bool:
        return (fuentes / fichero).exists()

    poblacion: dict[str, int] = {}  # de Wikidata (P1082)

    # --- Cataluña (Municat): ayuntamientos ---
    provincia_comarca: dict[str, Counter] = defaultdict(Counter)  # para ubicar los consejos comarcales
    if existe("cat-ajuntaments.csv"):
        for r in leer_csv(fuentes / "cat-ajuntaments.csv"):
            provincia_comarca[normalizar(r.get("comarca", ""))][r["codi_ens"][:2]] += 1
            anotar(r["codi_ens"][:5], "cat-ajuntaments.csv", telefono=telefono(r.get("telefon")), email=emails(r.get("e_mail")),
                   web=web(r.get("web")), alcalde=nombre_propio(r.get("president", "")), partido=r.get("partitpol", ""),
                   direccion=f'{r.get("adre_a", "")}, {r.get("codipostal", "")} {r.get("municipi", "")}'.strip(", "))
    # --- Euskadi ---
    if existe("euskadi-entidades.json"):
        for r in json.loads((fuentes / "euskadi-entidades.json").read_text(encoding="utf-8")):
            if r.get("entType") != "Municipio":
                continue
            anotar(r["territorycode"] + r["municipalitycode"], "euskadi-entidades.json", telefono=telefono(r.get("phone")),
                   email=emails(r.get("email")), web=web(r.get("webpage")), alcalde=nombre_propio(r.get("entMajorName", "")),
                   partido=r.get("entPoliticalParty", ""), direccion=f'{r.get("address", "").strip()}, {r.get("postalcode", "")} {r.get("municipality", "")}')
    # --- Galicia ---
    if existe("galicia-concellos.csv"):
        for r in leer_csv(fuentes / "galicia-concellos.csv"):
            v = list(r.values())  # CONCELLO, PROVINCIA, ENDEREZO, CP, TELÉFONO, FAX, CORREO, WEB...
            anotar(municipios.buscar(v[0], codigo_provincia(v[1])), "galicia-concellos.csv", telefono=telefono(v[4]),
                   email=emails(v[6]), web=web(v[7]), direccion=f"{v[2]}, {v[3]} {v[0]}")
    # --- Castilla-La Mancha ---
    if existe("clm-ayuntamientos.csv"):
        for r in leer_csv(fuentes / "clm-ayuntamientos.csv"):
            v = list(r.values())  # ENTIDAD_LOCAL;EATIM;MUNICIPIO;PROVINCIA;DIRECCION;CP;TELEFONO;CORREO_E;WEB;ALCALDE/SA;PARTIDO
            if v[1].strip().lower() != "no":
                continue
            anotar(municipios.buscar(v[2], codigo_provincia(v[3])), "clm-ayuntamientos.csv", telefono=telefono(v[6]),
                   email=emails(v[7]), web=web(v[8]), alcalde=nombre_propio(v[9]), partido=v[10], direccion=f"{v[4]}, {v[5]} {v[2]}")
    # --- Comunitat Valenciana: Castellón (Diputación) y toda la comunidad (Generalitat) ---
    if existe("castellon-ayuntamientos.csv"):
        for r in leer_csv(fuentes / "castellon-ayuntamientos.csv"):
            if r.get("id", r.get("tipo", "")) != "Ayuntamiento":
                continue
            cod = str(r["ine"]).split(".")[0].zfill(5)
            alcalde = re.sub(r"^(D\.|Dña\.|Dª|D\.ª|Da\.)\s*", "", r.get("alcalde", "").strip())
            anotar(cod, "castellon-ayuntamientos.csv", telefono=telefono(r.get("telefono")), email=emails(r.get("correo_electronico_ayuntamiento")),
                   web=web(r.get("web_ayuntamiento")), alcalde=alcalde, partido=r.get("partido", ""),
                   direccion=f'{r.get("domicilio", "")}, {str(r.get("cod_postal", "")).split(".")[0]} {r.get("nombre_poblacion", "")}')
    if existe("gva-entidades-locales.csv"):
        for r in leer_csv(fuentes / "gva-entidades-locales.csv"):
            cod = re.sub(r"\D", "", str(r.get("codigo_ine", ""))).zfill(5)
            anotar(cod, "gva-entidades-locales.csv", telefono=telefono(r.get("telefono")), email=emails(r.get("mail")),
                   direccion=f'{r.get("direccion", "")}, {r.get("cp", "")} {r.get("nombre", "")}')
    # --- Navarra ---
    if existe("navarra-entidades.csv"):
        for r in leer_csv(fuentes / "navarra-entidades.csv"):
            if not r.get("Tipo entidad (ES)", "").startswith("Municipio"):
                continue
            anotar(municipios.buscar(r["Denominación (ES)"], "31"), "navarra-entidades.csv", telefono=telefono(r.get("Teléfono")),
                   email=emails(r.get("Email")), web=web(r.get("Web")), alcalde=nombre_propio(r.get("Titular", "")), direccion=r.get("Dirección", ""))
    # --- Andalucía (RAEL) ---
    if existe("andalucia-rael.txt"):
        for r in leer_csv(fuentes / "andalucia-rael.txt", "|"):
            if r.get("claseEntidad") != "MUNICIPIO":
                continue
            try:
                reg = ast.literal_eval(r["registrales"]) if r.get("registrales") else {}
                capital = ast.literal_eval(reg["capitalidad"]) if isinstance(reg.get("capitalidad"), str) else reg.get("capitalidad") or {}
                pres = ast.literal_eval(r["presidente"]) if r.get("presidente") else {}
                dire = ast.literal_eval(r["direccion"]) if r.get("direccion") else {}
            except (ValueError, SyntaxError):
                sin_ubicar["andalucia-rael.txt"] += 1
                continue
            alcalde = " ".join(x for x in (pres.get("nombre"), pres.get("apellido1"), pres.get("apellido2")) if x)
            via = " ".join(x.strip() for x in (dire.get("tipoVia") or "", dire.get("nombreVia") or "", dire.get("numeroVia") or "") if x and x.strip())
            anotar(municipios.buscar(reg.get("denominacion", ""), codigo_provincia(capital.get("provincia", ""))), "andalucia-rael.txt",
                   telefono=telefono(r.get("telefono")), email=emails(r.get("correo")), web=web(r.get("direccionWeb")),
                   alcalde=nombre_propio(alcalde), partido=pres.get("grupoPolitico") or "",
                   direccion=f'{nombre_propio(via)}, {dire.get("codPostal", "")} {nombre_propio(dire.get("municipio", ""))}')
    # --- Aragón ---
    if existe("aragon-municipios.csv"):
        for r in leer_csv(fuentes / "aragon-municipios.csv"):
            anotar(str(r["codigo_mun"]).zfill(5), "aragon-municipios.csv", telefono=telefono(r.get("telefono")), email=emails(r.get("email")),
                   alcalde=nombre_propio(r.get("alcalde", "")), direccion=f'{nombre_propio(r.get("direccion", ""))}, {r.get("cp", "")} {r.get("nombre", "")}')
    # --- Salamanca (alcaldes) ---
    if existe("salamanca-alcaldes.csv"):
        for r in leer_csv(fuentes / "salamanca-alcaldes.csv"):
            if not str(r.get("CODIGO", "")).strip().isdigit():
                continue
            nombre = " ".join(x.strip() for x in (r.get("NOMBRE"), r.get("PRIMER APELLIDO"), r.get("SEGUNDO APELLIDO")) if x and x.strip())
            anotar("37" + r["CODIGO"].strip().zfill(3), "salamanca-alcaldes.csv", alcalde=nombre_propio(nombre), partido=r.get("PARTIDO", ""))
    # --- Wikidata (solo la web, para los municipios sin otra fuente) ---
    if existe("wikidata-webs.csv"):
        for r in leer_csv(fuentes / "wikidata-webs.csv", ","):
            anotar(r["ine"], "wikidata-webs.csv", web=web(r.get("web")))
            if re.fullmatch(r"\d+(\.0+)?", r.get("poblacion") or ""):
                poblacion[r["ine"]] = int(float(r["poblacion"]))

    # --- Cargos con área de cultura ---
    cultura: list[dict] = []
    vistos = set()
    if existe("cat-carrecs.csv"):
        for r in leer_csv(fuentes / "cat-carrecs.csv"):
            if not CULTURA.search(r.get("area") or ""):
                continue
            ens = r["nom_complert"]
            es_ajuntament = ens.startswith("Ajuntament")
            cod = r["codi_ens"][:5] if es_ajuntament else None
            nombre = nombre_propio(r["nom_regidor"])
            if cod:
                vistos.add((cod, tokens_persona(nombre)))
            cpro = cod[:2] if cod else None
            comarca = re.match(r"Consell Comarcal (?:de les |dels |del |de la |de l'|d'|de )?(.*)", ens)
            if comarca and provincia_comarca.get(normalizar(comarca[1])):
                cpro = provincia_comarca[normalizar(comarca[1])].most_common(1)[0][0]
            elif ens.startswith("Diputació de "):
                cpro = codigo_provincia(ens.removeprefix("Diputació de "))
            cultura.append({"codigo": cod, "cpro": cpro, "entidad": ens, "tipo": tipo_entidad(ens),
                            "nombre": nombre, "cargo": r.get("carrec", ""), "area": r.get("area", ""), "partido": r.get("partit", ""),
                            "desde": (r.get("data_nomenament") or "")[:10], "fuente": "cat-carrecs.csv"})
    if existe("diba-electes.json"):
        for r in json.loads((fuentes / "diba-electes.json").read_text(encoding="utf-8")).get("elements", []):
            if not CULTURA.search(r.get("ambit_treball") or ""):
                continue
            mun = (r.get("codi_ine_persona") or [{}])[0]
            nombre = nombre_propio(" ".join(x for x in (r.get("nom"), r.get("cognom1_persona"), r.get("cognom2_persona")) if x))
            nom = mun.get("municipi_nom", "")
            ens = f"Ajuntament d'{nom}" if re.match(r"[AEIOUHÀÈÉÍÒÓÚ]", nom) else f"Ajuntament de {nom}"
            # La misma persona ya viene de Municat (con el nombre escrito de otra forma): se queda la de Municat.
            if any(cod == mun.get("ine") and misma_persona(t, tokens_persona(nombre)) for cod, t in vistos):
                continue
            cultura.append({"codigo": mun.get("ine"), "cpro": "08", "entidad": ens, "tipo": "Ayuntamiento", "nombre": nombre,
                            "cargo": r.get("carrec_persona", ""), "area": r.get("ambit_treball", ""),
                            "partido": r.get("sigles_partit_politic", ""), "desde": "", "fuente": "diba-electes.json"})

    cultura_por_municipio: dict[str, list[dict]] = defaultdict(list)
    for c in cultura:
        if c["codigo"]:
            cultura_por_municipio[c["codigo"]].append(c)

    # --- Tablas de salida ---
    filas_ayto = []
    for m in sorted(municipios.lista, key=lambda m: tuple(normalizar(x) for x in (*PROVINCIAS[m["codigo"][:2]][::-1], m["nombre"]))):
        cod = m["codigo"]
        prov, ccaa = PROVINCIAS[cod[:2]]
        f = ficha.get(cod, {})
        resp = cultura_por_municipio.get(cod, [])
        origen = sorted({descripcion.get(fich, fich) for _, fich in f.values()} | {descripcion.get(c["fuente"], "") for c in resp} - {""})
        filas_ayto.append({
            "Comunidad autónoma": ccaa, "Provincia": prov, "Municipio": m["nombre"], "Código INE": cod,
            "Población": poblacion.get(cod, ""),
            "Responsable de cultura": "; ".join(c["nombre"] for c in resp),
            "Cargo y área": "; ".join(f'{c["cargo"]}: {c["area"]}' for c in resp),
            "Teléfono ayuntamiento": f.get("telefono", ("",))[0], "Email ayuntamiento": f.get("email", ("",))[0],
            "Web": f.get("web", ("",))[0], "Alcalde/sa": f.get("alcalde", ("",))[0], "Partido alcaldía": f.get("partido", ("",))[0],
            "Dirección": f.get("direccion", ("",))[0], "Fuentes": "; ".join(origen),
        })
    contacto_por_codigo = {f["Código INE"]: f for f in filas_ayto}
    filas_cultura = []
    for c in cultura:
        f = contacto_por_codigo.get(c["codigo"] or "", {})
        prov, ccaa = PROVINCIAS.get(c["cpro"] or "", ("", "Cataluña"))
        filas_cultura.append({
            "Comunidad autónoma": ccaa, "Provincia": prov, "Entidad": c["entidad"], "Tipo de entidad": c["tipo"],
            "Municipio": f.get("Municipio", ""), "Código INE": c["codigo"] or "", "Población": f.get("Población", ""),
            "Nombre": c["nombre"], "Cargo": c["cargo"], "Área de gobierno": c["area"], "Partido": c["partido"],
            "Nombramiento": c["desde"], "Teléfono ayuntamiento": f.get("Teléfono ayuntamiento", ""),
            "Email ayuntamiento": f.get("Email ayuntamiento", ""), "Web": f.get("Web", ""), "Fuente": descripcion.get(c["fuente"], c["fuente"]),
        })
    filas_cultura.sort(key=lambda r: (normalizar(r["Provincia"]), r["Tipo de entidad"] != "Ayuntamiento",
                                      normalizar(r["Municipio"] or r["Entidad"]), normalizar(r["Nombre"])))

    resumen = []
    grupos: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for f in filas_ayto:
        grupos[(f["Comunidad autónoma"], f["Provincia"])].append(f)
    for (ccaa, prov), fs in sorted(grupos.items(), key=lambda g: (normalizar(g[0][0]), normalizar(g[0][1]))):
        resumen.append({"Comunidad autónoma": ccaa, "Provincia": prov, "Municipios": len(fs),
                        "Con responsable de cultura": sum(1 for f in fs if f["Responsable de cultura"]),
                        "Con teléfono": sum(1 for f in fs if f["Teléfono ayuntamiento"]),
                        "Con email": sum(1 for f in fs if f["Email ayuntamiento"]),
                        "Con web": sum(1 for f in fs if f["Web"]),
                        "Con alcalde/sa": sum(1 for f in fs if f["Alcalde/sa"])})
    total = {k: (sum(r[k] for r in resumen) if isinstance(resumen[0][k], int) else "") for k in resumen[0]}
    total["Comunidad autónoma"] = "TOTAL ESPAÑA"
    resumen.append(total)

    escribir_csv(salida / "ayuntamientos.csv", filas_ayto)
    escribir_csv(salida / "cultura.csv", filas_cultura)
    escribir_excel(salida / "directorio-cultura.xlsx", filas_cultura, filas_ayto, resumen, estado)

    print(f"Municipios: {len(filas_ayto)} · cargos de cultura: {len(filas_cultura)}")
    for k in ("Con responsable de cultura", "Con teléfono", "Con email", "Con web", "Con alcalde/sa"):
        print(f"  {k}: {total[k]}")
    if sin_ubicar:
        print("Registros sin ubicar por fuente:", dict(sin_ubicar))


def escribir_csv(ruta: Path, filas: list[dict]) -> None:
    with ruta.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(filas[0].keys()))
        w.writeheader()
        w.writerows(filas)


def escribir_excel(ruta: Path, cultura: list[dict], aytos: list[dict], resumen: list[dict], estado: dict) -> None:
    libro = openpyxl.Workbook()
    cabecera = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="1F4E79")

    def hoja(titulo: str, filas: list[dict], anchos: dict[str, int], primera: bool = False):
        ws = libro.active if primera else libro.create_sheet()
        ws.title = titulo
        columnas = list(filas[0].keys())
        ws.append(columnas)
        for f in filas:
            ws.append([f[c] for c in columnas])
        for i, c in enumerate(columnas, 1):
            celda = ws.cell(row=1, column=i)
            celda.font, celda.fill = cabecera, relleno
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = anchos.get(c, 16)
            if c == "Web":
                for fila in range(2, len(filas) + 2):
                    v = ws.cell(row=fila, column=i).value
                    if v:
                        ws.cell(row=fila, column=i).hyperlink = v
                        ws.cell(row=fila, column=i).style = "Hyperlink"
            if c == "Población":
                for fila in range(2, len(filas) + 2):
                    ws.cell(row=fila, column=i).number_format = "#,##0"
        ws.freeze_panes = "D2" if titulo != "Resumen" else "C2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32
        return ws

    anchos = {"Comunidad autónoma": 20, "Provincia": 16, "Municipio": 24, "Entidad": 34, "Nombre": 28, "Cargo": 20,
              "Área de gobierno": 50, "Responsable de cultura": 28, "Cargo y área": 50, "Email ayuntamiento": 30, "Web": 30,
              "Alcalde/sa": 28, "Dirección": 36, "Fuentes": 40, "Fuente": 30, "Teléfono ayuntamiento": 16, "Código INE": 9}
    hoja("Cultura", cultura, anchos, primera=True)
    hoja("Ayuntamientos", aytos, anchos)
    ws = hoja("Resumen", resumen, {"Comunidad autónoma": 26, "Provincia": 24})
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)

    notas = libro.create_sheet("Fuentes y notas")
    notas.column_dimensions["A"].width = 120
    lineas = [
        ("Directorio de ayuntamientos y responsables de cultura", True),
        (f"Fuentes descargadas el {estado['fecha']}. Se actualiza cada mes con el workflow «Directorio de cultura».", False),
        ("", False),
        ("Qué hay y qué no", True),
        ("· No existe ninguna lista oficial de concejales de cultura de toda España. Solo Cataluña publica el área de gobierno de cada concejal "
         "(Municat, Generalitat; y la Diputació de Barcelona para su provincia). Por eso la hoja «Cultura» solo tiene cargos de Cataluña.", False),
        ("· Para el resto de España, la hoja «Ayuntamientos» da el contacto general del ayuntamiento (teléfono, email, web) y el alcalde/sa "
         "cuando la comunidad o la diputación lo publica en abierto. Desde ahí se pide «la concejalía de cultura».", False),
        ("· Comunidades con directorio oficial completo: Cataluña, País Vasco (web y alcalde), Galicia (sin alcalde), Castilla-La Mancha, "
         "Comunitat Valenciana (sin alcalde salvo Castellón), Navarra, Andalucía y Aragón. Salamanca: solo alcaldes. "
         "Resto: solo la web del ayuntamiento (Wikidata).", False),
        ("· La población es la última cifra de Wikidata (padrón del INE), solo orientativa para ordenar.", False),
        ("· Los teléfonos y emails son los institucionales del ayuntamiento, no de la persona. Los cargos cambian: comprobad antes de escribir.", False),
        ("· Uso: son datos públicos. Para enviar comunicaciones comerciales por email, tened en cuenta la LSSI y el RGPD "
         "(identificaros, explicar de dónde sale el contacto y ofrecer la baja).", False),
        ("", False),
        ("Fuentes (todas con licencia abierta; citar la fuente al reutilizar)", True),
    ]
    for f in estado["fuentes"]:
        lineas.append((f"· {f['descripcion']} — {f['licencia']} — {f['url']} — {f['resultado']}", False))
    for i, (texto, negrita) in enumerate(lineas, 1):
        c = notas.cell(row=i, column=1, value=texto)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if negrita:
            c.font = Font(bold=True, size=12)
    libro.save(ruta)


if __name__ == "__main__":
    main()
